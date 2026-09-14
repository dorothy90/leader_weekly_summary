"""Word export and uploaded reference parsing for Group Report Studio."""

from __future__ import annotations

import io
from pathlib import Path
import unicodedata
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.shared import Mm


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ZIP_FILES = 200
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
KOREAN_FONT = "Nanum Gothic"


def _set_font(style, name: str, size: float | None = None) -> None:
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        style.font.size = Pt(size)


def _remove_paragraph_borders(paragraph_or_style) -> None:
    element = paragraph_or_style._element
    p_pr = element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margin(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def _prevent_row_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _display_width(value: str) -> int:
    return sum(2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1 for char in value)


def _set_table_borders(table, color: str = "D9D9D9") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def _set_table_grid_widths(table, widths: list[int]) -> None:
    grid_cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
    for grid_col, width in zip(grid_cols, widths):
        grid_col.set(qn("w:w"), str(round(width / 635)))


def _canonicalize_docx(raw: bytes) -> bytes:
    """Remove ZIP timestamp variance while preserving package entry order."""
    source = zipfile.ZipFile(io.BytesIO(raw), "r")
    output = io.BytesIO()
    with source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for old in source.infolist():
            info = zipfile.ZipInfo(old.filename, _FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = old.external_attr
            info.create_system = 0
            target.writestr(info, source.read(old.filename))
    return output.getvalue()


def export_docx(report: dict) -> bytes:
    """Render an HTTP-schema report dictionary as deterministic DOCX bytes."""
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    _set_font(styles["Normal"], KOREAN_FONT, 11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for style_name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13)):
        style = styles[style_name]
        _set_font(style, KOREAN_FONT, size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.font.underline = False
        _remove_paragraph_borders(style)

    title = str(report.get("title") or "주간 보고서")
    title_paragraph = doc.add_paragraph(title, style="Title")
    _remove_paragraph_borders(title_paragraph)
    for run in title_paragraph.runs:
        run.font.underline = False
    week = str(report.get("week") or "-")
    version = str(report.get("version", 0))
    metadata = doc.add_paragraph(f"주차 {week}  버전 {version}")
    metadata.alignment = WD_ALIGN_PARAGRAPH.LEFT
    metadata.runs[0].font.color.rgb = RGBColor(99, 120, 138)
    metadata.paragraph_format.space_after = Pt(14)

    last_group = object()
    for section_data in report.get("sections") or []:
        group = section_data.get("group")
        if group not in (None, "") and group != last_group:
            doc.add_paragraph(str(group), style="Heading 1")
            last_group = group

        section_title = section_data.get("title")
        if section_title:
            doc.add_paragraph(str(section_title), style="Heading 2")

        for block in section_data.get("blocks") or []:
            kind = block.get("kind")
            if kind == "paragraph":
                doc.add_paragraph(str(block.get("text") or ""))
            elif kind == "bullet":
                doc.add_paragraph(str(block.get("text") or ""), style="List Bullet")
            elif kind == "table":
                headers = [str(value) for value in (block.get("headers") or [])]
                rows = [[str(value) for value in row] for row in (block.get("rows") or [])]
                column_count = max([len(headers), *(len(row) for row in rows)], default=0)
                if not column_count:
                    continue
                table = doc.add_table(rows=1 if headers else 0, cols=column_count)
                table.style = "Table Grid"
                table.autofit = False
                usable_width = section.page_width - section.left_margin - section.right_margin
                weights = []
                for index in range(column_count):
                    values = ([headers[index]] if index < len(headers) else []) + [
                        row[index] for row in rows if index < len(row)
                    ]
                    weights.append(max(6, min(36, max((_display_width(value) for value in values), default=6))))
                total_weight = sum(weights)
                widths = [round(usable_width * weight / total_weight) for weight in weights]
                widths[-1] += usable_width - sum(widths)

                _set_table_borders(table)
                _set_table_grid_widths(table, widths)

                if headers:
                    header_row = table.rows[0]
                    _repeat_table_header(header_row)
                    _prevent_row_split(header_row)
                    for index, cell in enumerate(header_row.cells):
                        cell.text = headers[index] if index < len(headers) else ""
                        cell.width = widths[index]
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        _set_cell_margin(cell)
                        _shade(cell, "183B56")
                        for run in cell.paragraphs[0].runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                for row_number, values in enumerate(rows):
                    row = table.add_row()
                    _prevent_row_split(row)
                    cells = row.cells
                    for index, cell in enumerate(cells):
                        cell.text = values[index] if index < len(values) else ""
                        cell.width = widths[index]
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        _set_cell_margin(cell)
                        if row_number % 2:
                            _shade(cell, "F2F6F9")
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            else:
                raise ValueError(f"지원하지 않는 블록 형식입니다: {kind}")

    stream = io.BytesIO()
    doc.save(stream)
    return _canonicalize_docx(stream.getvalue())


def _validate_docx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content), "r") as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_FILES:
                raise ValueError("DOCX 내부 파일 수가 허용 범위를 초과합니다.")
            total_size = sum(member.file_size for member in members)
            if total_size > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("DOCX 압축 해제 크기가 허용 범위를 초과합니다.")
            if "word/document.xml" not in archive.namelist():
                raise ValueError("올바른 DOCX 파일이 아닙니다.")
    except ValueError:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise ValueError("올바른 DOCX 파일이 아닙니다.") from exc


def _docx_text(content: bytes) -> str:
    _validate_docx_archive(content)
    try:
        doc = Document(io.BytesIO(content))
    except Exception as exc:
        raise ValueError("올바른 DOCX 파일이 아닙니다.") from exc

    lines: list[str] = []
    body = doc.element.body
    def node_text(node) -> str:
        return "".join(text_node.text or "" for text_node in node.xpath(".//w:t"))

    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            text = node_text(child).strip()
            if text:
                lines.append(text)
        elif child.tag == qn("w:tbl"):
            for tr in child.findall(qn("w:tr")):
                cells = []
                for tc in tr.findall(qn("w:tc")):
                    paragraphs = [node_text(p).strip() for p in tc.findall(qn("w:p"))]
                    cells.append(" ".join(text for text in paragraphs if text))
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def parse_reference(name: str, content: bytes) -> str:
    """Extract ordered plain text from an uploaded DOCX, UTF-8 TXT, or Markdown file."""
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("업로드 파일 크기가 허용 범위를 초과합니다.")
    suffix = Path(name).suffix.lower()
    if suffix in {".txt", ".md"}:
        try:
            return content.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ValueError("TXT와 Markdown 파일은 UTF-8 인코딩이어야 합니다.") from exc
    if suffix == ".docx":
        return _docx_text(content)
    raise ValueError("지원하지 않는 파일 형식입니다. DOCX, TXT, MD 파일을 사용하세요.")
