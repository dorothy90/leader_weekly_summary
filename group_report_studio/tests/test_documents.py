import io
import unittest
import zipfile

from docx import Document
from docx.shared import Mm

from studio.documents import export_docx, parse_reference


class DocumentExportTests(unittest.TestCase):
    def test_export_preserves_section_and_block_order(self):
        report = {
            "title": "주간 그룹 보고",
            "week": "2026-W37",
            "version": 3,
            "sections": [
                {
                    "id": "events",
                    "group": "0. 핵심 이벤트",
                    "title": "이번 주 핵심",
                    "blocks": [
                        {"kind": "paragraph", "text": "첫 문단"},
                        {"kind": "bullet", "text": "첫 항목"},
                        {
                            "kind": "table",
                            "headers": ["팀", "내용"],
                            "rows": [["A", "완료"], ["B", "진행"]],
                        },
                        {"kind": "paragraph", "text": "마지막 문단"},
                    ],
                },
                {
                    "id": "plans",
                    "group": "1. 다음 주 계획",
                    "title": "계획",
                    "blocks": [{"kind": "paragraph", "text": "다음 작업"}],
                },
            ],
        }

        raw = export_docx(report)
        doc = Document(io.BytesIO(raw))
        body_items = []
        for child in doc.element.body.iterchildren():
            if child.tag.endswith("}p"):
                text = "".join(node.text or "" for node in child.xpath(".//w:t"))
                if text:
                    body_items.append(("p", text))
            elif child.tag.endswith("}tbl"):
                rows = []
                for tr in child.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr"):
                    rows.append(["".join(node.text or "" for node in tc.xpath(".//w:t")) for tc in tr.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc")])
                body_items.append(("table", rows))

        self.assertEqual(body_items[0], ("p", "주간 그룹 보고"))
        self.assertIn(("p", "주차 2026-W37  버전 3"), body_items)
        expected = [
            ("p", "0. 핵심 이벤트"),
            ("p", "이번 주 핵심"),
            ("p", "첫 문단"),
            ("p", "첫 항목"),
            ("table", [["팀", "내용"], ["A", "완료"], ["B", "진행"]]),
            ("p", "마지막 문단"),
            ("p", "1. 다음 주 계획"),
            ("p", "계획"),
            ("p", "다음 작업"),
        ]
        start = body_items.index(expected[0])
        self.assertEqual(body_items[start : start + len(expected)], expected)

        styles = doc.styles
        self.assertEqual(styles["Title"].font.color.rgb.__str__(), "000000")
        self.assertFalse(styles["Title"].font.underline)
        self.assertFalse(styles["Title"]._element.xpath("./w:pPr/w:pBdr"))
        self.assertFalse(doc.paragraphs[0]._p.xpath("./w:pPr/w:pBdr"))
        self.assertEqual(styles["Normal"]._element.rPr.rFonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"), "Nanum Gothic")
        header_markers = doc.tables[0].rows[0]._tr.xpath("./w:trPr/w:tblHeader")
        self.assertEqual(len(header_markers), 1)
        self.assertEqual(header_markers[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"), "true")
        self.assertAlmostEqual(doc.sections[0].page_width / Mm(1), 210, places=0)
        self.assertAlmostEqual(doc.sections[0].page_height / Mm(1), 297, places=0)
        borders = doc.tables[0]._tbl.xpath("./w:tblPr/w:tblBorders/*")
        self.assertEqual(len(borders), 6)
        self.assertTrue(all(border.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color") == "D9D9D9" for border in borders))
        self.assertTrue(all(row._tr.xpath("./w:trPr/w:cantSplit") for row in doc.tables[0].rows))

    def test_export_is_deterministic(self):
        report = {"title": "보고서", "week": "2026-W37", "version": 1, "sections": []}
        self.assertEqual(export_docx(report), export_docx(report))


class ReferenceParsingTests(unittest.TestCase):
    def test_txt_and_markdown_are_utf8(self):
        self.assertEqual(parse_reference("지난주.txt", "안녕하세요".encode()), "안녕하세요")
        self.assertEqual(parse_reference("notes.MD", b"# Heading\nBody"), "# Heading\nBody")

    def test_docx_preserves_paragraph_table_interleaving(self):
        doc = Document()
        doc.add_paragraph("앞 문단")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "항목"
        table.rows[0].cells[1].text = "값"
        doc.add_paragraph("뒤 문단")
        stream = io.BytesIO()
        doc.save(stream)

        self.assertEqual(parse_reference("prior.docx", stream.getvalue()), "앞 문단\n항목 | 값\n뒤 문단")

    def test_rejects_unsupported_invalid_and_oversized_content(self):
        with self.assertRaisesRegex(ValueError, "지원하지 않는 파일 형식"):
            parse_reference("data.pdf", b"pdf")
        with self.assertRaisesRegex(ValueError, "UTF-8"):
            parse_reference("bad.txt", b"\xff")
        with self.assertRaisesRegex(ValueError, "올바른 DOCX"):
            parse_reference("bad.docx", b"not a zip")
        with self.assertRaisesRegex(ValueError, "파일 크기"):
            parse_reference("large.md", b"x" * (10 * 1024 * 1024 + 1))

    def test_rejects_docx_zip_bomb_metadata(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            info = zipfile.ZipInfo("word/document.xml")
            info.file_size = 51 * 1024 * 1024
            archive.writestr(info, b"tiny")
        # A highly compressible payload reliably crosses the aggregate expanded-size limit.
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"x" * (51 * 1024 * 1024))
        with self.assertRaisesRegex(ValueError, "압축 해제"):
            parse_reference("bomb.docx", stream.getvalue())


if __name__ == "__main__":
    unittest.main()
