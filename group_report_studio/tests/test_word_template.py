import base64
import io
import tempfile
import unittest
from pathlib import Path

import httpx
from docx import Document

from studio.api import create_app
from studio.config import Settings
from studio.store import Store
from test_harness import FakeLLM, FakeSource


class TemplateLLM(FakeLLM):
    def __init__(self):
        super().__init__()
        self.template_calls = []
        self.writing_calls = []

    def complete(self, stage, system, payload, schema):
        if stage == 'word_template':
            self.template_calls.append(payload)
            return dict(name='간결한 주보', writing_prompt='소주제당 글머리표 2개 이내로 작성한다.',
                        sections=[dict(id='spica',group='1. 수율',title='Spica',instructions='')])
        if stage in ('write','edit'):
            self.writing_calls.append(payload)
        return super().complete(stage, system, payload, schema)


class WordTemplateTests(unittest.IsolatedAsyncioTestCase):
    async def test_extract_review_save_default_and_generate_without_references(self):
        with tempfile.TemporaryDirectory() as folder:
            llm = TemplateLLM()
            app = create_app(Settings(data_dir=Path(folder)),source=FakeSource(),llm=llm,inline=True)
            self.addCleanup(app.state.executor.shutdown, wait=True)
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://test') as client:
                doc = Document()
                doc.add_heading('1. 수율', level=1)
                doc.add_paragraph('Spica 수율 91.2%, 검사 완료.')
                buf = io.BytesIO()
                doc.save(buf)
                response = await client.post('/api/templates/from-word',json=dict(name='주보.docx',content_base64=base64.b64encode(buf.getvalue()).decode()))
                self.assertEqual(response.status_code,200,response.text)
                draft = response.json()
                self.assertIn('91.2%', llm.template_calls[0]['text'])
                self.assertNotIn('91.2%',draft['writing_prompt'])
                self.assertEqual((await client.get('/api/templates')).json(), [])
                draft['writing_prompt'] = '소주제당 2문장 이내.'
                saved = await client.post('/api/templates/default',json=draft)
                self.assertEqual(saved.status_code,200,saved.text)
                self.assertEqual(Store(Path(folder)/'studio.db').default_template()['writing_prompt'],draft['writing_prompt'])
                config = (await client.get('/api/config')).json()
                self.assertEqual(config['default_template'],draft)
                report = (await client.post('/api/reports',json=dict(week='2026-01',template=draft))).json()
                job = (await client.post('/api/reports/'+report['id']+'/generate',json=dict(base_version=0))).json()
                self.assertEqual((await client.get('/api/jobs/'+job['id'])).json()['status'],'succeeded')
                self.assertEqual(report['references'],[])
                self.assertEqual(llm.writing_calls[0]['writing_prompt'],draft['writing_prompt'])
                await client.post('/api/reports/'+report['id']+'/edit',json=dict(base_version=1,message='줄여줘',section_id='spica'))
                self.assertEqual(llm.writing_calls[-1]['writing_prompt'],draft['writing_prompt'])
                changed = dict(draft,writing_prompt='새 기본 규칙')
                await client.post('/api/templates/default',json=changed)
                existing = (await client.get('/api/reports/'+report['id'])).json()
                self.assertEqual(existing['template']['writing_prompt'],draft['writing_prompt'])

    async def test_invalid_files_do_not_call_model(self):
        with tempfile.TemporaryDirectory() as folder:
            llm = TemplateLLM()
            app = create_app(Settings(data_dir=Path(folder)),source=FakeSource(),llm=llm,inline=True)
            self.addCleanup(app.state.executor.shutdown, wait=True)
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://test') as client:
                for name, content in [('a.txt',b'hi'),('a.docx',b'not a zip')]:
                    response = await client.post('/api/templates/from-word',json=dict(name=name,content_base64=base64.b64encode(content).decode()))
                    self.assertEqual(response.status_code,400,response.text)
                self.assertFalse(llm.template_calls)

    async def test_empty_and_oversized_documents_are_not_silently_truncated(self):
        with tempfile.TemporaryDirectory() as folder:
            llm = TemplateLLM()
            app = create_app(Settings(data_dir=Path(folder)),source=FakeSource(),llm=llm,inline=True)
            self.addCleanup(app.state.executor.shutdown, wait=True)
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://test') as client:
                for text in ['', '긴 본문입니다. ' * 10000]:
                    doc = Document()
                    doc.add_paragraph(text)
                    buf = io.BytesIO()
                    doc.save(buf)
                    response = await client.post('/api/templates/from-word',json=dict(name='주보.docx',content_base64=base64.b64encode(buf.getvalue()).decode()))
                    self.assertEqual(response.status_code,400,response.text)
                self.assertFalse(llm.template_calls)
